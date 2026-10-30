# --- Log, notifications and the 4 tools the AI can use ---
import base64                                       # encodes images and PDFs for the API
import filecmp                                      # compares two files byte by byte
import io                                           # files in memory
import os                                           # folders and environment variables
import re                                           # splits names into words
import shutil                                       # moves files
import subprocess                                   # runs PowerShell for notifications
from datetime import datetime                       # date and time for the log
from pathlib import Path                            # path handling

from anthropic import beta_tool                     # turns a Python function into a tool for the AI
from docx import Document                           # reads Word files
from openpyxl import load_workbook                  # reads Excel files
from PIL import Image                               # shrinks images before sending
from pptx import Presentation                       # reads PowerPoint files
from pypdf import PdfReader, PdfWriter              # reads PDFs, extracts a page

from config import DEPTH, DL, HOME, IMAGES, LOG, ROOTS, TO_SORT


# --- Log: written to the screen AND to the file ---
def log(message):
    line = f"{datetime.now():%Y-%m-%d %H:%M:%S} | {message}"
    print(line)
    with open(LOG, "a", encoding="utf-8") as f:                        # "a" = append without erasing
        f.write(line + "\n")


# --- Windows notification: shows where a file went ---
TOAST = (
    "[Windows.UI.Notifications.ToastNotificationManager, Windows.UI.Notifications, ContentType = WindowsRuntime] | Out-Null;"
    "$x = [Windows.UI.Notifications.ToastNotificationManager]::GetTemplateContent([Windows.UI.Notifications.ToastTemplateType]::ToastText02);"
    "$t = $x.GetElementsByTagName('text');"
    "$t.Item(0).AppendChild($x.CreateTextNode($env:TOAST_TITLE)) | Out-Null;"
    "$t.Item(1).AppendChild($x.CreateTextNode($env:TOAST_TEXT)) | Out-Null;"
    "if ($env:TOAST_LINK) { $x.DocumentElement.SetAttribute('activationType', 'protocol');"  # click = opens the link
    " $x.DocumentElement.SetAttribute('launch', $env:TOAST_LINK) };"
    "[Windows.UI.Notifications.ToastNotificationManager]::CreateToastNotifier("
    "'{1AC14E77-02E7-4E5D-B744-2EB1AE5198B7}\\WindowsPowerShell\\v1.0\\powershell.exe')"
    ".Show([Windows.UI.Notifications.ToastNotification]::new($x))")


def notify(title, text, link=None):                                     # link: folder or file opened by a click
    try:                                                                # texts passed as variables: no injection through a file name
        subprocess.Popen(["powershell.exe", "-NoProfile", "-Command", TOAST],
                         env={**os.environ, "TOAST_TITLE": title, "TOAST_TEXT": text,
                              "TOAST_LINK": link.as_uri() if link else ""},
                         creationflags=subprocess.CREATE_NO_WINDOW)     # no window, no waiting
    except Exception as e:
        log(f"  notification failed: {e}")


# --- Security: the AI can only touch files placed directly in Downloads ---
def in_downloads(name):
    f = (DL / name).resolve()
    if f.parent != DL.resolve() or not f.is_file():                     # blocks "../Documents/..." and folders
        log(f"  refused: {name} (outside Downloads)")
        return None
    return f


def as_base64(data):
    return base64.standard_b64encode(data).decode()


# --- Tool 1: see the existing folders ---
@beta_tool
def list_folders() -> str:
    """Lists the existing folders and subfolders where a file can be sorted."""
    folders = []
    for root in ROOTS:
        for dirpath, subdirs, _ in os.walk(root):
            subdirs[:] = sorted(s for s in subdirs if not s.startswith("."))  # skips hidden folders
            if len(Path(dirpath).relative_to(root).parts) >= DEPTH:
                subdirs.clear()                                         # does not go deeper
            folders.append(Path(dirpath).relative_to(HOME).as_posix())  # e.g. Documents/Bank/Statements
    return "\n".join(folders[:500])                                     # capped to stay fast


# --- Tool 2: find similar files already sorted ---
@beta_tool
def find_similar(name: str) -> str:
    """Searches the sorted folders for files whose name resembles this one (same subject, another version, same series).
    A similar file already sorted often points to the right folder."""
    words = {w for w in re.split(r"[\W_]+", Path(name).stem.lower())   # words of the name, accents included
             if len(w) >= 4 and not w.isdigit()}                        # skips "ia", "(1)", "2026"…
    if not words:
        return "Name too short to search for similar files"
    found = []
    for root in ROOTS:
        for dirpath, subdirs, files in os.walk(root):
            subdirs[:] = [s for s in subdirs if not s.startswith(".")]
            for f in files:
                common = words & set(re.split(r"[\W_]+", Path(f).stem.lower()))
                if len(common) >= min(2, len(words)):                   # at least 2 words in common
                    found.append((len(common), (Path(dirpath) / f).relative_to(HOME).as_posix()))
    if not found:
        return "No similar file already sorted"
    found.sort(reverse=True)                                            # most similar first
    return "\n".join(f"{path} ({n} words in common)" for n, path in found[:10])


# --- Tool 3: read the beginning of a file ---
@beta_tool
def read_start(name: str) -> str | list:
    """Reads the beginning of a PDF, Word, Excel, PowerPoint, text or image file from the Downloads folder.
    For an image or a scanned PDF, returns the image of the first page."""
    f = in_downloads(name)
    if not f:
        return "File not found in Downloads"
    ext = f.suffix.lower()
    try:
        if ext == ".pdf":
            pdf = PdfReader(f)
            text = pdf.pages[0].extract_text() or ""                    # first page
            if len(text.strip()) < 20:                                  # no text: scanned PDF
                page = PdfWriter()
                page.add_page(pdf.pages[0])                             # first page only, to limit cost
                buffer = io.BytesIO()
                page.write(buffer)
                log(f"  {name}: scanned PDF, first page sent to the AI")
                return [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                        "data": as_base64(buffer.getvalue())}},
                        {"type": "text", "text": "Scanned PDF: here is its first page."}]
        elif ext in IMAGES:
            img = Image.open(f)
            img.seek(0)                                                 # first frame of an animated GIF or multipage TIFF
            img = img.convert("RGB")
            img.thumbnail((1568, 1568))                                 # shrunk: fewer tokens, still readable
            buffer = io.BytesIO()
            img.save(buffer, "JPEG", quality=85)
            log(f"  {name}: image sent to the AI")
            return [{"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                                                 "data": as_base64(buffer.getvalue())}},
                    {"type": "text", "text": "Here is the image."}]
        elif ext == ".docx":
            text = "\n".join(p.text for p in Document(f).paragraphs)    # paragraphs
        elif ext in (".xlsx", ".xlsm"):
            sheet = load_workbook(f, read_only=True, data_only=True).active
            text = "\n".join(" | ".join(str(c) for c in row if c is not None)
                             for row in sheet.iter_rows(max_row=20, values_only=True))  # 20 rows
        elif ext == ".pptx":
            text = "\n".join(shape.text_frame.text                      # slide text, in order
                             for slide in Presentation(f).slides
                             for shape in slide.shapes if shape.has_text_frame)
        else:
            with open(f, "rb") as h:
                head = h.read(8000)                                     # 8 KB only, even for a 2 GB video
            if b"\x00" in head:                                         # null bytes = binary (image, program, archive…)
                log(f"  {name}: binary file, not read")
                return "Binary file (image, program, archive…): content unreadable, decide from the name."
            text = head.decode("utf-8", errors="ignore")                # plain text
        log(f"  reading the content of {name}")
        return text[:1500]                                              # 1500 characters max
    except Exception as e:
        return f"Unreadable: {e}"


# --- Tool 4: move a file ---
@beta_tool
def move_file(name: str, destination: str, reason: str) -> str:
    """Moves a file to an existing folder (path as given by list_folders) or to 'To_sort'.
    reason: explain in one short sentence why you chose this folder."""
    src = in_downloads(name)
    if not src:                                                         # already moved or outside Downloads
        return "File not found in Downloads"
    name = src.name
    if destination == "To_sort":
        folder = TO_SORT
        folder.mkdir(exist_ok=True)                                     # the only folder it can create
    else:
        folder = (HOME / destination).resolve()
        allowed = any(folder == r or r in folder.parents for r in ROOTS)
        if not allowed or not folder.is_dir():                          # security: existing and allowed only
            log(f"  refused: {destination}")
            return "Destination not allowed or missing, pick one from the list"
    target = folder / name
    if target.is_file() and filecmp.cmp(src, target, shallow=False):    # same name AND same content, byte by byte
        src.unlink()                                                    # useless copy: the original is already sorted
        log(f"DUPLICATE | {name} already in {folder} (identical) -> downloaded copy deleted")
        notify(f"Duplicate deleted: {name}", f"Already in {folder.relative_to(HOME).as_posix()}", folder)
        return "already there, identical: copy deleted"
    if target.exists():                                                 # same name, different content: avoid overwriting
        target = folder / f"{src.stem}_{datetime.now():%Y-%m-%d_%H%M%S}{src.suffix}"
    shutil.move(src, target)                                            # actual move
    log(f"SORTED | {name} -> {target} | reason: {reason}")
    if folder == TO_SORT:
        notify(f"To sort: {name}", reason, folder)                      # the user has to decide
    else:
        notify(f"Sorted: {target.name}", folder.relative_to(HOME).as_posix(), folder)  # the move is too fast to see: shows where it went
    return "moved"
