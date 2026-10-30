# --- Libraries ---
import os, re, io, time, base64, shutil, filecmp, subprocess, anthropic  # processes, words, memory, time, encoding, files, notifications, Claude API
from datetime import datetime                       # date and time for the log
from pathlib import Path                            # path handling
from anthropic import beta_tool                     # turns a Python function into a tool for the AI
from pypdf import PdfReader, PdfWriter              # reads PDFs, extracts a page
from PIL import Image                               # shrinks images before sending
from docx import Document                           # reads Word files
from openpyxl import load_workbook                  # reads Excel files
from pptx import Presentation                       # reads PowerPoint files
from watchdog.observers import Observer             # watches a folder
from watchdog.events import FileSystemEventHandler  # reacts to new files

# --- Configuration ---
UTILISATEUR = "Phil"                                # first name used in the prompt: change it
LANGUE = "French"                                   # language of the reasons (log, notifications): change it
HOME = Path.home()                                  # e.g. C:\Users\<name>
DL = HOME / "Downloads"                             # watched folder
A_TRIER = DL / "A_trier"                            # destination when in doubt
RACINES = [HOME / "Documents", HOME / "Pictures",   # only allowed places (and their subfolders)
           HOME / "Music", HOME / "Videos"]
JOURNAL = Path(__file__).parent / "journal.log"     # log file, next to the script
RATTRAPAGE_JOURS = 7                                # on startup, sorts files less than 7 days old
PID = Path(__file__).parent / "agent.pid"           # process ID, to stop only this agent
PROFONDEUR = 4                                      # subfolder levels shown to the AI (Documents/Work/Client/Project/2026)
IMAGES = (".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tif", ".tiff")  # shown to the AI, which can read them

# --- Log: written to the screen AND to the file ---
def journal(message):
    ligne = f"{datetime.now():%Y-%m-%d %H:%M:%S} | {message}"
    print(ligne)
    with open(JOURNAL, "a", encoding="utf-8") as f:                    # "a" = append without erasing
        f.write(ligne + "\n")

# --- Windows notification: warns when a file is not sorted ---
TOAST = (
    "[Windows.UI.Notifications.ToastNotificationManager, Windows.UI.Notifications, ContentType = WindowsRuntime] | Out-Null;"
    "$x = [Windows.UI.Notifications.ToastNotificationManager]::GetTemplateContent([Windows.UI.Notifications.ToastTemplateType]::ToastText02);"
    "$t = $x.GetElementsByTagName('text');"
    "$t.Item(0).AppendChild($x.CreateTextNode($env:TOAST_TITRE)) | Out-Null;"
    "$t.Item(1).AppendChild($x.CreateTextNode($env:TOAST_TEXTE)) | Out-Null;"
    "[Windows.UI.Notifications.ToastNotificationManager]::CreateToastNotifier("
    "'{1AC14E77-02E7-4E5D-B744-2EB1AE5198B7}\\WindowsPowerShell\\v1.0\\powershell.exe')"
    ".Show([Windows.UI.Notifications.ToastNotification]::new($x))")

def notifier(titre, texte):
    try:                                                                # texts passed as variables: no injection through a file name
        subprocess.Popen(["powershell.exe", "-NoProfile", "-Command", TOAST],
                         env={**os.environ, "TOAST_TITRE": titre, "TOAST_TEXTE": texte},
                         creationflags=subprocess.CREATE_NO_WINDOW)     # no window, no waiting
    except Exception as e:
        journal(f"  notification impossible : {e}")

# --- Security: the AI can only touch files placed directly in Downloads ---
def dans_telechargements(nom):
    f = (DL / nom).resolve()
    if f.parent != DL.resolve() or not f.is_file():                     # blocks "../Documents/..." and folders
        journal(f"  refusé : {nom} (hors Téléchargements)")
        return None
    return f

# --- Tool 1: see the existing folders ---
@beta_tool
def lister_dossiers() -> str:
    """Lists the existing folders and subfolders where a file can be sorted."""
    dossiers = []
    for racine in RACINES:
        for chemin, sous, _ in os.walk(racine):
            sous[:] = sorted(s for s in sous if not s.startswith("."))  # skips hidden folders
            if len(Path(chemin).relative_to(racine).parts) >= PROFONDEUR:
                sous.clear()                                            # does not go deeper
            dossiers.append(Path(chemin).relative_to(HOME).as_posix())  # e.g. Documents/Bank/Statements
    return "\n".join(dossiers[:500])                                    # capped to stay fast

# --- Tool 2: find similar files already sorted ---
@beta_tool
def chercher_similaires(nom: str) -> str:
    """Searches the sorted folders for files whose name resembles this one (same subject, another version, same series).
    A similar file already sorted often points to the right folder."""
    mots = {m for m in re.split(r"[\W_]+", Path(nom).stem.lower())     # words of the name, accents included
            if len(m) >= 4 and not m.isdigit()}                         # skips "ia", "(1)", "2026"…
    if not mots:
        return "Nom trop court pour chercher des fichiers proches"
    trouves = []
    for racine in RACINES:
        for chemin, sous, fichiers in os.walk(racine):
            sous[:] = [s for s in sous if not s.startswith(".")]
            for f in fichiers:
                communs = mots & set(re.split(r"[\W_]+", Path(f).stem.lower()))
                if len(communs) >= min(2, len(mots)):                   # at least 2 words in common
                    trouves.append((len(communs), (Path(chemin) / f).relative_to(HOME).as_posix()))
    if not trouves:
        return "Aucun fichier proche déjà rangé"
    trouves.sort(reverse=True)                                          # most similar first
    return "\n".join(f"{chemin} ({n} mots communs)" for n, chemin in trouves[:10])

# --- Tool 3: read the beginning of a file ---
@beta_tool
def lire_debut(nom: str) -> str | list:
    """Reads the beginning of a PDF, Word, Excel, PowerPoint, text or image file from the Downloads folder.
    For an image or a scanned PDF, returns the image of the first page."""
    f = dans_telechargements(nom)
    if not f:
        return "Fichier introuvable dans Téléchargements"
    ext = f.suffix.lower()
    try:
        if ext == ".pdf":
            pdf = PdfReader(f)
            texte = pdf.pages[0].extract_text() or ""                   # first page
            if len(texte.strip()) < 20:                                 # no text: scanned PDF
                page = PdfWriter(); page.add_page(pdf.pages[0])         # first page only, to limit cost
                tampon = io.BytesIO(); page.write(tampon)
                journal(f"  {nom} : PDF scanné, 1re page envoyée à l'IA")
                return [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                         "data": base64.standard_b64encode(tampon.getvalue()).decode()}},
                        {"type": "text", "text": "PDF scanné : voici sa 1re page."}]
        elif ext in IMAGES:
            img = Image.open(f); img.seek(0)                            # first frame of an animated GIF or multipage TIFF
            img = img.convert("RGB"); img.thumbnail((1568, 1568))       # shrunk: fewer tokens, still readable
            tampon = io.BytesIO(); img.save(tampon, "JPEG", quality=85)
            journal(f"  {nom} : image envoyée à l'IA")
            return [{"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                     "data": base64.standard_b64encode(tampon.getvalue()).decode()}},
                    {"type": "text", "text": "Voici l'image."}]
        elif ext == ".docx":
            texte = "\n".join(p.text for p in Document(f).paragraphs)   # paragraphs
        elif ext in (".xlsx", ".xlsm"):
            feuille = load_workbook(f, read_only=True, data_only=True).active
            texte = "\n".join(" | ".join(str(c) for c in row if c is not None)
                              for row in feuille.iter_rows(max_row=20, values_only=True))  # 20 rows
        elif ext == ".pptx":
            texte = "\n".join(forme.text_frame.text                     # slide text, in order
                              for diapo in Presentation(f).slides
                              for forme in diapo.shapes if forme.has_text_frame)
        else:
            with open(f, "rb") as h:
                debut = h.read(8000)                                    # 8 KB only, even for a 2 GB video
            if b"\x00" in debut:                                        # null bytes = binary (image, program, archive…)
                journal(f"  {nom} : fichier binaire, non lu")
                return "Fichier binaire (image, programme, archive…) : contenu illisible, décide d'après le nom."
            texte = debut.decode("utf-8", errors="ignore")              # plain text
        journal(f"  lecture du contenu de {nom}")
        return texte[:1500]                                             # 1500 characters max
    except Exception as e:
        return f"Illisible : {e}"

# --- Tool 4: move a file ---
@beta_tool
def deplacer(nom: str, destination: str, raison: str) -> str:
    """Moves a file to an existing folder (path as given by lister_dossiers) or to 'A_trier'.
    raison: explain in one short sentence why you chose this folder."""
    src = dans_telechargements(nom)
    if not src:                                                         # already moved or outside Downloads
        return "Fichier introuvable dans Téléchargements"
    nom = src.name
    if destination == "A_trier":
        dossier = A_TRIER
        dossier.mkdir(exist_ok=True)                                    # the only folder it can create
    else:
        dossier = (HOME / destination).resolve()
        autorise = any(dossier == r or r in dossier.parents for r in RACINES)
        if not autorise or not dossier.is_dir():                        # security: existing and allowed only
            journal(f"  refusé : {destination}")
            return "Destination non autorisée ou inexistante, choisis-en une de la liste"
    cible = dossier / nom
    if cible.is_file() and filecmp.cmp(src, cible, shallow=False):      # same name AND same content, byte by byte
        src.unlink()                                                    # useless copy: the original is already sorted
        journal(f"DOUBLON | {nom} déjà présent dans {dossier} (identique) -> copie téléchargée supprimée")
        return "déjà présent à l'identique, copie supprimée"
    if cible.exists():                                                  # same name, different content: avoid overwriting
        cible = dossier / f"{src.stem}_{datetime.now():%Y-%m-%d_%H%M%S}{src.suffix}"
    shutil.move(src, cible)                                             # actual move
    journal(f"RANGÉ   | {nom} -> {cible} | raison : {raison}")
    if dossier == A_TRIER:
        notifier(f"À trier : {nom}", raison)                            # the user has to decide
    return "déplacé"

# --- The agent: the AI decides (ROCIF structured prompt) ---
def ranger(nom):
    journal(f"NOUVEAU | {nom}")
    fin = anthropic.Anthropic(api_key=os.environ["DOC_SORTER_API_KEY"]).beta.messages.tool_runner(  # key dedicated to this agent
        model="claude-haiku-5-5", max_tokens=2000,                      # fast, cheap model
        tools=[lister_dossiers, chercher_similaires, lire_debut, deplacer],  # its 4 tools
        messages=[{"role": "user", "content":
            f"ROLE: You are {UTILISATEUR}'s file sorting assistant.\n"
            f"OBJECTIVE: Move the downloaded file \"{nom}\" to the right place.\n"
            f"CONTEXT: {UTILISATEUR} organizes files by organization or project (e.g. Documents/Bank) "
            "before organizing them by type. Specific subfolders are always preferred over generic folders.\n"
            "INSTRUCTIONS:\n"
            "1. List the existing folders and search for already sorted files with a similar name.\n"
            "2. If a similar file covers the same subject (another version, same series), put it in the same folder.\n"
            "3. Otherwise, if the name is not enough, read the beginning of the file.\n"
            "4. Put it in the most specific existing folder.\n"
            "5. If there is no folder for the organization or project, use a folder by document type "
            "if one exists (e.g. Documents/Invoices for an invoice). Put it in A_trier only if no "
            "folder fits, not even by type. Never guess.\n"
            f"FORMAT: A single move action, with a one-sentence reason written in {LANGUE}."}]
    ).until_done()                                                      # runs until the end
    if (DL / nom).exists():                                             # the AI stopped without moving the file
        texte = " ".join(b.text for b in fin.content if b.type == "text").replace("\n", " ")
        journal(f"NON RANGÉ | {nom} | arrêt : {fin.stop_reason} | réponse : {texte[:300]}")
        deplacer(nom, "A_trier", "l'IA n'a pas choisi de dossier")      # safety net

# --- Waits until the file is fully written (stable size) ---
def attendre_fin(p, max_secondes=600):
    taille = -1
    for _ in range(max_secondes):
        time.sleep(1)
        if not p.is_file():
            return False                                                # disappeared in the meantime
        if (DL / (p.name + ".part")).exists():                          # Firefox: empty file + .part in progress
            return False                                                # we'll be called again when the .part is renamed
        nouvelle = p.stat().st_size
        if nouvelle == taille:
            return True                                                 # 2 identical readings: download finished
        taille = nouvelle
    journal(f"ERREUR  | {p.name} | taille encore instable après {max_secondes} s")
    notifier(f"Erreur : {p.name}", "Téléchargement jamais terminé, fichier non rangé")
    return False

# --- The trigger ---
class Handler(FileSystemEventHandler):
    def on_created(self, e): self.check(e.src_path)                     # new file
    def on_moved(self, e): self.check(e.dest_path)                      # renamed file (end of download)
    def check(self, p):
        p = Path(p)
        try:
            if p.parent == DL and p.suffix.lower() not in (".crdownload", ".tmp", ".part"):  # skips downloads in progress
                if attendre_fin(p):
                    ranger(p.name)                                      # runs the agent
        except Exception as e:
            journal(f"ERREUR  | {p.name} | {e}")                        # errors go to the log too
            notifier(f"Erreur : {p.name}", str(e)[:200])                # e.g. expired API key

# --- Startup ---
obs = Observer(); obs.schedule(Handler(), str(DL)); obs.start()         # watches Downloads
PID.write_text(str(os.getpid()))                                        # allows a targeted stop
journal(f"DÉMARRAGE | surveillance de {DL} | PID {os.getpid()}")

# --- Catch-up: files that arrived while the agent was stopped ---
def deja_vus():
    """Last time the agent processed each file name, according to the log."""
    vus = {}
    if JOURNAL.exists():
        for ligne in JOURNAL.read_text(encoding="utf-8", errors="ignore").splitlines():
            date, _, reste = ligne.partition(" | ")
            if reste.startswith("NOUVEAU | "):
                try:
                    vus[reste[10:]] = datetime.strptime(date, "%Y-%m-%d %H:%M:%S").timestamp()
                except ValueError:
                    pass                                                # damaged line: skipped
    return vus

limite = time.time() - RATTRAPAGE_JOURS * 86400
vus = deja_vus()
recents = [p for p in sorted(DL.iterdir()) if p.is_file() and p.stat().st_mtime > limite]
en_attente = [p for p in recents if p.stat().st_mtime > vus.get(p.name, 0)]  # already processed and unchanged since: left on purpose
journal(f"RATTRAPAGE | {len(en_attente)} fichier(s) de moins de {RATTRAPAGE_JOURS} jours"
        f" ({len(recents) - len(en_attente)} déjà traité(s), laissé(s) en place)")
for p in en_attente:
    Handler().check(p)                                                  # same filters as a new download

while True: time.sleep(1)                                               # Ctrl + C to stop
